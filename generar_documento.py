# -*- coding: utf-8 -*-
"""Genera el documento Word de la Práctica 6 con capturas y respuestas."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).resolve().parent
CAP = BASE / "capturas"
OUT = BASE / "Practica6_MobileConnect.docx"


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_heading_custom(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=16 if level == 1 else 13, bold=True, color=RGBColor(0x1B, 0x4F, 0x72))


def add_para(doc: Document, text: str, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=11, bold=bold)
    p.paragraph_format.space_after = Pt(8)
    return p


def add_image(doc: Document, path: Path, width_cm: float = 8.0, caption: str = ""):
    if not path.exists():
        add_para(doc, f"[Falta imagen: {path.name}]")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    if caption:
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.add_run(caption)
        set_run_font(r, size=10, color=RGBColor(0x55, 0x55, 0x55))


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Práctica 6: Interacción entre dispositivos móviles")
    set_run_font(r, size=18, bold=True, color=RGBColor(0x1B, 0x4F, 0x72))

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("MobileConnect — WebSocket (contrato v2) · sala practica6")
    set_run_font(r, size=12, color=RGBColor(0x33, 0x33, 0x33))

    add_para(
        doc,
        "Objetivo: demostrar comunicación real entre dos teléfonos mediante un servidor "
        "intermediario WebSocket (canal_servidor). El servidor asigna los roles A/B; los "
        "clientes se unen a la misma sala, intercambian mensajes, muestran hora y contadores, "
        "y validan texto vacío antes de enviar.",
    )
    add_para(
        doc,
        "Canal usado: ws://10.13.8.252:8080/ws · sala: practica6. "
        "Primer dispositivo que envía texto válido = A; quien recibe = B.",
    )

    # --- Evidencias propias ---
    add_heading_custom(doc, "1. Evidencias — teléfono propio (cliente MobileConnect)", 1)

    add_heading_custom(doc, "1.1 Chat conectado (rol B)", 2)
    add_para(
        doc,
        "Conexión activa al servidor de la pareja. Estado «Conectado · sala practica6 · emisor B». "
        "Se observa intercambio: mensaje de A («hola») y respuesta de B («qpdo»), con hora y contadores.",
    )
    add_image(doc, CAP / "01_mio_chat_conectado.png", 7.5, "Figura 1. Teléfono propio: chat en red como Dispositivo B.")

    add_heading_custom(doc, "1.2 Mensajes A y B + contadores", 2)
    add_para(
        doc,
        "Historial con burbujas de distinto color según emisor. Contadores Msgs A / Msgs B "
        "actualizados según la lista (estado Compose).",
    )
    add_image(doc, CAP / "02_mio_mensajes_ab.png", 7.5, "Figura 2. Teléfono propio: varios mensajes A/B y contadores.")

    add_heading_custom(doc, "1.3 Desconexión / error de red", 2)
    add_para(
        doc,
        "Evidencia de comunicación real: al cerrarse el socket aparece «Error de conexión: Socket closed». "
        "En una simulación local no existiría este fallo de red.",
    )
    add_image(doc, CAP / "03_mio_socket_cerrado.png", 7.5, "Figura 3. Teléfono propio: caída del WebSocket.")

    add_heading_custom(doc, "1.4 Validación / estado antes de enviar", 2)
    add_para(
        doc,
        "La app avisa al usuario (snackbar «Conéctate primero») y no envía basura. "
        "También se valida trim() para no mandar mensajes vacíos.",
    )
    add_image(doc, CAP / "04_mio_validacion_snackbar.png", 7.5, "Figura 4. Teléfono propio: validación de estado al enviar.")

    # --- Pareja ---
    add_heading_custom(doc, "2. Evidencias — teléfono de la pareja (servidor / cliente)", 1)

    add_heading_custom(doc, "2.1 Pantalla de conexión al canal", 2)
    add_para(
        doc,
        "Configuración de IP del servidor, puerto 8080 y sala practica6. "
        "Se explica que el servidor asigna roles (primer envío = A; quien recibe = B).",
    )
    add_image(doc, CAP / "08_pareja_pantalla_conectar.png", 7.5, "Figura 5. Pareja: pantalla Conectar (IP / puerto / sala).")

    add_heading_custom(doc, "2.2 Sala lista sin rol aún", 2)
    add_para(
        doc,
        "Tras conectar, estado «En red» pero «Sin rol» hasta el primer mensaje válido (regla R1 del contrato v2).",
    )
    add_image(doc, CAP / "07_pareja_sin_rol.png", 7.5, "Figura 6. Pareja: conectado, esperando asignación A/B.")

    add_heading_custom(doc, "2.3 Primer intercambio", 2)
    add_para(
        doc,
        "Chat en vivo: Dispositivo A y B con «hola», horas ISO convertidas a hora local y contadores A/B.",
    )
    add_image(doc, CAP / "06_pareja_hola.png", 7.5, "Figura 7. Pareja: mensajes bidireccionales.")

    add_heading_custom(doc, "2.4 Conversación identificada", 2)
    add_para(
        doc,
        "Sesión completa en sala practica6 · 10.13.8.252. Este teléfono como Dispositivo B; "
        "mensajes «soy angel» / «soy Pablo» confirman dos clientes reales distintos.",
    )
    add_image(doc, CAP / "05_pareja_chat_angel_pablo.png", 7.5, "Figura 8. Pareja: chat A/B con identidades.")

    # --- Condición switch ---
    add_heading_custom(doc, "3. Condición switch / case (roles A y B)", 1)
    add_para(
        doc,
        "En el contrato WebSocket los emisores viajan como cadenas «A» / «B». "
        "El código convierte el enum interno a código wire y viceversa con un switch; "
        "si llega un valor inválido lanza FormatException.",
    )
    add_image(doc, CAP / "09_switch_codigo_wire.png", 14.0, "Figura 9. Switch: codigoWire / fromCodigoWire (A ↔ B).")

    # --- Preguntas ---
    add_heading_custom(doc, "4. Respuestas a las preguntas", 1)

    qa = [
        (
            "4.1 ¿Qué significa la interacción entre dispositivos móviles?",
            "Es el intercambio de información entre dos o más teléfonos de forma que una acción "
            "en uno se refleja en el otro. En esta práctica: un mensaje enviado desde un celular "
            "aparece en el del compañero a través de un canal WebSocket compartido (misma sala en el servidor).",
        ),
        (
            "4.2 ¿Qué tecnologías permiten intercambiar información entre dispositivos?",
            "Bluetooth / BLE, Wi‑Fi Direct, Wi‑Fi/LAN, internet (HTTP/WebSocket), NFC y APIs como "
            "Nearby Connections. En MobileConnect usamos WebSocket sobre Wi‑Fi/LAN "
            "(ws://IP:8080/ws) hacia el servidor canal_servidor de la pareja.",
        ),
        (
            "4.3 ¿Qué diferencia existe entre una simulación local y una comunicación real?",
            "En la simulación ambos «dispositivos» viven en la misma app/memoria (modo local A/B). "
            "En la comunicación real hay dos procesos en dos teléfonos; los datos cruzan la red "
            "(IP del servidor, sala practica6), con latencia, posibles desconexiones (Socket closed) "
            "y roles asignados por el servidor.",
        ),
        (
            "4.4 ¿Qué función cumple el estado de la aplicación en Jetpack Compose?",
            "Guarda lo que la UI muestra (mensajes, conexión, rol A/B, contadores, texto del campo). "
            "Al cambiar el estado, Compose recompone y actualiza la pantalla. Sin estado reactivo, "
            "el chat no se refrescaría al recibir un mensaje o una asignación por WebSocket.",
        ),
        (
            "4.5 ¿Por qué es importante validar los mensajes antes de enviarlos?",
            "Evita basura en el historial, reduce tráfico inútil por la red y mejora la experiencia "
            "(no enviar cadenas vacías o solo espacios). El cliente hace trim(); el servidor también "
            "rechaza texto vacío con un error.",
        ),
        (
            "4.6 ¿Qué cambios serían necesarios para implementar comunicación real mediante Bluetooth o Wi‑Fi?",
            "Permisos (Bluetooth/ubicación/Wi‑Fi), descubrimiento o emparejamiento, sockets o una API "
            "(Nearby, BluetoothSocket, Wi‑Fi Direct), corrutinas, serializar mensajes y manejar caídas "
            "de enlace. En esta práctica la comunicación real se resolvió con Wi‑Fi + WebSocket y un "
            "servidor intermediario, sin emparejar Bluetooth punto a punto.",
        ),
        (
            "4.7 ¿Qué ventajas tendría utilizar un servidor intermediario para intercambiar mensajes?",
            "Los teléfonos no necesitan emparejarse entre sí: solo se conectan al mismo canal. El "
            "servidor guarda historial, asigna roles A/B, hace broadcast y puede escalar a más clientes. "
            "Facilita depuración (/salud) y reinicio con limpiar. En internet permitiría chat a "
            "distancia, autenticación y notificaciones.",
        ),
    ]
    for heading, body in qa:
        add_heading_custom(doc, heading, 2)
        add_para(doc, body)

    add_heading_custom(doc, "5. Resumen técnico", 1)
    for item in [
        "Cliente Kotlin/Compose (MobileConnect) + servidor Dart WebSocket (canal_servidor).",
        "Contrato v2: unirse(sala), enviar(texto), limpiar; recibe historial, mensaje, asignacion, error.",
        "Roles A/B asignados por el servidor (no los elige el cliente).",
        "Validación de mensaje vacío (cliente + servidor).",
        "Contadores por emisor A/B y hora desde enviadoEn (ISO‑8601).",
        "Repositorio: https://github.com/AngelSanchezVerduzco/practica6",
    ]:
        p = doc.add_paragraph(item, style="List Bullet")
        for run in p.runs:
            set_run_font(run, size=11)

    doc.save(OUT)
    print(f"OK: {OUT}")


if __name__ == "__main__":
    main()
